"""
SafeZoo AI 통합 안전관제 시스템 - PPT 자동 생성 및 편집 에이전트 스크립트
(Auto-Generation & Revision Pipeline for SafeZoo Presentation)
"""

import io
import sys
from pathlib import Path

# Windows cp949 콘솔 출력을 위한 UTF-8 인코딩 설정
if sys.stdout.encoding != "utf-8":
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ---------------------------------------------------------------------------
# 테마 및 컬러 시스템 (Dark & Red Accent Theme)
# ---------------------------------------------------------------------------
COLOR_BG = RGBColor(0x14, 0x14, 0x14)       # #141414 (배경 다크)
COLOR_CARD = RGBColor(0x23, 0x23, 0x23)     # #232323 (카드 블록)
COLOR_CARD_BORDER = RGBColor(0x38, 0x38, 0x38)
COLOR_TEXT_MAIN = RGBColor(0xFF, 0xFF, 0xFF)  # 흰색
COLOR_TEXT_MUTED = RGBColor(0xB3, 0xB3, 0xB3) # 연회색
COLOR_RED = RGBColor(0xE5, 0x09, 0x14)       # #E50914 (안전 경고 레드)
COLOR_GREEN = RGBColor(0x00, 0xC8, 0x53)     # #00C853 (SAFE 초록)
COLOR_BLUE_ACCENT = RGBColor(0x29, 0xB6, 0xF6)

BASE_DIR = Path(__file__).resolve().parent
OUTPUT_PPTX = BASE_DIR / "SafeZoo_Presentation.pptx"


def set_slide_background(slide, color=COLOR_BG):
    """슬라이드 전체 배경색 설정"""
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_header(slide, title_text, category_text="SAFEZOO CONTROL SYSTEM"):
    """공통 상단 헤더 생성"""
    # 카테고리 태그
    txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(10), Inches(0.4))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = category_text.upper()
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = COLOR_RED

    # 메인 제목
    txBox2 = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.5), Inches(0.8))
    tf2 = txBox2.text_frame
    tf2.word_wrap = True
    p2 = tf2.paragraphs[0]
    p2.text = title_text
    p2.font.size = Pt(24)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_TEXT_MAIN


def add_card(slide, left, top, width, height, bg_color=COLOR_CARD, border_color=COLOR_CARD_BORDER):
    """카드 스타일 사각형 배경 생성"""
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    return shape


# ---------------------------------------------------------------------------
# 슬라이드별 생성기
# ---------------------------------------------------------------------------
def build_slide_1(prs):
    """Slide 01: 표지"""
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # blank layout
    set_slide_background(slide)

    # 상단 장식 배지
    add_card(slide, Inches(0.8), Inches(1.5), Inches(3.2), Inches(0.4), bg_color=COLOR_RED, border_color=None)
    tx_badge = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(3.2), Inches(0.4))
    p = tx_badge.text_frame.paragraphs[0]
    p.text = "AI VISION SAFETY SYSTEM"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEXT_MAIN
    p.alignment = PP_ALIGN.CENTER

    # 타이틀
    tx = slide.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(11.5), Inches(1.8))
    tf = tx.text_frame
    tf.word_wrap = True
    p1 = tf.paragraphs[0]
    p1.text = "SafeZoo: AI 기반 동물원 통합 안전관제"
    p1.font.size = Pt(38)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_TEXT_MAIN

    p2 = tf.add_paragraph()
    p2.text = "YOLO 객체 추적 & ROI 경계 감지 기반 실시간 위험 대응 플랫폼"
    p2.font.size = Pt(20)
    p2.font.color.rgb = COLOR_TEXT_MUTED
    p2.space_before = Pt(12)

    # 하단 메타 정보 카드
    add_card(slide, Inches(0.8), Inches(5.0), Inches(11.7), Inches(1.5))
    tx_meta = slide.shapes.add_textbox(Inches(1.1), Inches(5.2), Inches(11.1), Inches(1.1))
    tf_meta = tx_meta.text_frame
    
    p = tf_meta.paragraphs[0]
    p.text = "핵심 모듈: zoo_safety.py (비전 엔진)  |  zoo_vision.py (Gradio 관제 대시보드)"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEXT_MAIN

    p_sub = tf_meta.add_paragraph()
    p_sub.text = "주요 기능: Real-time Object Tracking • Safe/Danger Polygon ROI • Event Log CSV • Snapshot • Control Center Inbox"
    p_sub.font.size = Pt(12)
    p_sub.font.color.rgb = COLOR_TEXT_MUTED
    p_sub.space_before = Pt(8)


def build_slide_2(prs):
    """Slide 02: 문제 제기 & 도입 배경"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_header(slide, "동물원 안전 관리의 도전과제와 AI 관제 도입 필요성")

    # 좌측: 기존 방식의 한계 (Card)
    add_card(slide, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
    tx = slide.shapes.add_textbox(Inches(1.1), Inches(2.1), Inches(5.0), Inches(4.2))
    tf = tx.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "⚠️ 기존 수동 관제 방식의 한계"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_RED

    bullets = [
        "넓은 사육장 및 관람 구역으로 인한 모니터링 사각지대 존재",
        "다수 CCTV 화면 동시 감시 시 관제요원 피로도 증가 및 골든타임 누락",
        "동물 사육장 이탈 또는 관람객 위험구역 침범 시 즉각적 경고 부재",
        "사후 경위 파악을 위한 데이터 기록 및 타임라인 구성의 어려움"
    ]
    for b in bullets:
        p_b = tf.add_paragraph()
        p_b.text = "• " + b
        p_b.font.size = Pt(13)
        p_b.font.color.rgb = COLOR_TEXT_MUTED
        p_b.space_before = Pt(14)

    # 우측: SafeZoo 해결책 (Card)
    add_card(slide, Inches(6.9), Inches(1.8), Inches(5.6), Inches(4.8))
    tx2 = slide.shapes.add_textbox(Inches(7.2), Inches(2.1), Inches(5.0), Inches(4.2))
    tf2 = tx2.text_frame
    tf2.word_wrap = True
    
    p = tf2.paragraphs[0]
    p.text = "✅ SafeZoo AI 지능형 안전관제"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_GREEN

    bullets2 = [
        "YOLOv8/11 객체 탐지 알고리즘 기반 24시간 실시간 앵커 추적",
        "SAFE / DANGER 다각형(Polygon) 구역 자동 판정 시스템",
        "이탈 즉시 자동 스냅샷 저장 및 관제 센터 팝업/인박스 즉시 전송",
        "CSV 데이터 베이스 축적으로 사고 통계 분석 및 위험 지역 도출 지원"
    ]
    for b in bullets2:
        p_b = tf2.add_paragraph()
        p_b.text = "• " + b
        p_b.font.size = Pt(13)
        p_b.font.color.rgb = COLOR_TEXT_MAIN
        p_b.space_before = Pt(14)


def build_slide_3(prs):
    """Slide 03: 시스템 아키텍처"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_header(slide, "SafeZoo 시스템 전체 파이프라인 & 아키텍처")

    steps = [
        ("01. 비전 입력", "CCTV Video Stream\n(giraffe_cctv.mp4)", COLOR_BLUE_ACCENT),
        ("02. AI 탐지 엔진", "zoo_safety.py\nYOLO + Bottom Anchor", COLOR_RED),
        ("03. ROI 구역 판정", "SAFE vs DANGER\nPolygon Point Test", COLOR_GREEN),
        ("04. 이벤트 트리거", "Snapshot JPG / Event CSV\nControl Center Inbox", COLOR_BLUE_ACCENT),
        ("05. 통합 관제 대시보드", "zoo_vision.py\nGradio + 5s Auto Refresh", COLOR_RED),
    ]

    left_start = 0.8
    width = 2.15
    gap = 0.23

    for i, (title, desc, color) in enumerate(steps):
        cur_left = Inches(left_start + i * (width + gap))
        add_card(slide, cur_left, Inches(2.2), Inches(width), Inches(4.2))
        
        # 스텝 번호 바
        add_card(slide, cur_left, Inches(2.2), Inches(width), Inches(0.5), bg_color=color, border_color=None)
        tx_num = slide.shapes.add_textbox(cur_left, Inches(2.2), Inches(width), Inches(0.5))
        p_num = tx_num.text_frame.paragraphs[0]
        p_num.text = title
        p_num.font.size = Pt(12)
        p_num.font.bold = True
        p_num.font.color.rgb = COLOR_TEXT_MAIN
        p_num.alignment = PP_ALIGN.CENTER

        # 상세 내용
        tx_desc = slide.shapes.add_textbox(cur_left + Inches(0.1), Inches(3.0), Inches(width - 0.2), Inches(3.0))
        tf_d = tx_desc.text_frame
        tf_d.word_wrap = True
        p_d = tf_d.paragraphs[0]
        p_d.text = desc
        p_d.font.size = Pt(13)
        p_d.font.color.rgb = COLOR_TEXT_MAIN
        p_d.alignment = PP_ALIGN.CENTER


def build_slide_4(prs):
    """Slide 04: AI 탐지 엔진 & ROI 판정"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_header(slide, "핵심 기술 1: YOLO 객체 추적 및 다각형 ROI 경계 판정")

    # 좌측: 핵심 알고리즘 설명 카드
    add_card(slide, Inches(0.8), Inches(1.8), Inches(6.0), Inches(4.8))
    tx = slide.shapes.add_textbox(Inches(1.1), Inches(2.1), Inches(5.4), Inches(4.2))
    tf = tx.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "🔍 객체 추적 & 위치 판정 메커니즘"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_RED

    bullets = [
        "감시 클래스 지정: Giraffe(기린), Person(관람객/사육사) 모니터링",
        "발지점(Bottom Anchor) 좌표 추출:\nx = (x1 + x2)/2,  y = y2 (지면 접촉부)",
        "cv2.pointPolygonTest 다각형 테스트:\nSAFE_ZONE 내부에 발지점이 속하는지 실시간 판정",
        "Lost Frame 관리: 최대 20프레임 미탐지 시 트랙 메모리 자동 정돈(prune)"
    ]
    for b in bullets:
        p_b = tf.add_paragraph()
        p_b.text = "• " + b
        p_b.font.size = Pt(13)
        p_b.font.color.rgb = COLOR_TEXT_MAIN
        p_b.space_before = Pt(12)

    # 우측: 좌표 코드 스니펫 카드
    add_card(slide, Inches(7.1), Inches(1.8), Inches(5.4), Inches(4.8))
    tx2 = slide.shapes.add_textbox(Inches(7.4), Inches(2.1), Inches(4.8), Inches(4.2))
    tf2 = tx2.text_frame
    tf2.word_wrap = True

    p = tf2.paragraphs[0]
    p.text = "💻 zoo_safety.py 핵심 로직 예시"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_BLUE_ACCENT

    code_str = (
        "SAFE_ZONE = np.array([\n"
        "    [40, 300], [180, 230], [420, 200],\n"
        "    [700, 210], [830, 280], [820, 360],\n"
        "    [650, 420], [430, 440], ...\n"
        "])\n\n"
        "def zone_label(px, py, frame_h):\n"
        "    if in_poly(px, py, SAFE_ZONE):\n"
        "        return 'SAFE', COLOR_SAFE\n"
        "    return 'DANGER', COLOR_DANGER"
    )
    p_code = tf2.add_paragraph()
    p_code.text = code_str
    p_code.font.size = Pt(11)
    p_code.font.name = "Consolas"
    p_code.font.color.rgb = COLOR_TEXT_MUTED
    p_code.space_before = Pt(12)


def build_slide_5(prs):
    """Slide 05: 이벤트 제어 & 관제 알림 파이프라인"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_header(slide, "핵심 기술 2: 상태 전이 감지 및 3중 이벤트 알림 파이프라인")

    # 3개 이벤트 처리 출력 채널 카드
    channels = [
        ("📸 스냅샷 자동 저장", "snapshots/alert_YYYYMMDD_HHMMSS_id{N}.jpg\n\n위험 구역 진입 순간의 프레임을 고화질 JPG 파일로 즉시 캡처하여 증거 자료 보존", COLOR_RED),
        ("📊 CSV 이벤트 로그", "logs/zoo_events.csv\n\n시각, 카메라ID, TrackID, 객체, 구역, 이벤트유형(ZONE_EXIT), 심각도(HIGH), 신뢰도 수치 전파 저장", COLOR_BLUE_ACCENT),
        ("🚨 관제 센터 알림함", "logs/control_center_inbox.txt\n\n'위험구역을 벗어났으니 다른 CCTV를 확인하세요'\n실시간 인박스 텍스트 파일 갱신 및 터미널 출력", COLOR_GREEN),
    ]

    for i, (title, desc, color) in enumerate(channels):
        left = Inches(0.8 + i * 4.0)
        add_card(slide, left, Inches(2.0), Inches(3.7), Inches(4.5))

        tx = slide.shapes.add_textbox(left + Inches(0.2), Inches(2.2), Inches(3.3), Inches(4.0))
        tf = tx.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = color

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(13)
        p_d.font.color.rgb = COLOR_TEXT_MAIN
        p_d.space_before = Pt(14)


def build_slide_6(prs):
    """Slide 06: SafeZoo Vision 통합 관제 대시보드"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_header(slide, "SafeZoo Vision: Gradio 기반 실시간 통합 관제보드 (zoo_vision.py)")

    # 대시보드 구동 카드 1
    add_card(slide, Inches(0.8), Inches(1.8), Inches(7.5), Inches(4.8))
    tx = slide.shapes.add_textbox(Inches(1.1), Inches(2.1), Inches(6.9), Inches(4.2))
    tf = tx.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "🖥️ 대시보드 레이아웃 & UI 구성"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEXT_MAIN

    features = [
        "다크 & 레드 액센트 프리미엄 디자인 (safezoo.css 스타일 시트 적용)",
        "실시간 감시 탭: CCTV 비디오 플레이어 (giraffe_cctv.mp4) 자동 재생",
        "최근 알림 스냅샷 갤러리: 최신 6개 alert 스냅샷 모니터링",
        "관제센터 알림함: 실시간 메시지 수신 텍스트 상자",
        "gr.Timer(5초): 5초 주기 자동 데이터 갱신으로 백그라운드 탐지 결과 동기화"
    ]
    for f in features:
        p_f = tf.add_paragraph()
        p_f.text = "• " + f
        p_f.font.size = Pt(13)
        p_f.font.color.rgb = COLOR_TEXT_MUTED
        p_f.space_before = Pt(12)

    # 대시보드 테마 카드 2
    add_card(slide, Inches(8.6), Inches(1.8), Inches(3.9), Inches(4.8), bg_color=COLOR_RED, border_color=None)
    tx2 = slide.shapes.add_textbox(Inches(8.8), Inches(2.1), Inches(3.5), Inches(4.2))
    tf2 = tx2.text_frame
    tf2.word_wrap = True

    p2 = tf2.paragraphs[0]
    p2.text = "🔔 LIVE CONTROL"
    p2.font.size = Pt(20)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_TEXT_MAIN

    p2_desc = tf2.add_paragraph()
    p2_desc.text = "카메라 ID: ZOO_CAM_GIRAFFE_01\n상태: 🔴 ALERT / 🟢 NORMAL\n\n위험 감지 시 비상 팝업 배너가 상단에 노출되어 관제요원의 즉각적 대처 가능"
    p2_desc.font.size = Pt(13)
    p2_desc.font.color.rgb = COLOR_TEXT_MAIN
    p2_desc.space_before = Pt(16)


def build_slide_7(prs):
    """Slide 07: 통계 및 KPI 분석"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_header(slide, "통계 및 KPI 분석: 데이터 기반 지능형 데이터 시각화")

    # Top 3 KPI Cards
    kpis = [
        ("총 이벤트 건수", "COUNT", COLOR_TEXT_MAIN),
        ("HIGH 심각도 건수", "ALERTS", COLOR_RED),
        ("활성 카메라 수", "CAMERAS", COLOR_GREEN),
    ]
    for i, (title, sub, color) in enumerate(kpis):
        left = Inches(0.8 + i * 4.0)
        add_card(slide, left, Inches(1.8), Inches(3.7), Inches(1.5))
        tx = slide.shapes.add_textbox(left, Inches(2.0), Inches(3.7), Inches(1.1))
        tf = tx.text_frame
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(14)
        p.font.color.rgb = COLOR_TEXT_MUTED
        p.alignment = PP_ALIGN.CENTER
        
        p_val = tf.add_paragraph()
        p_val.text = sub
        p_val.font.size = Pt(22)
        p_val.font.bold = True
        p_val.font.color.rgb = color
        p_val.alignment = PP_ALIGN.CENTER

    # 하단 차트 설명 카드 2개
    add_card(slide, Inches(0.8), Inches(3.6), Inches(5.7), Inches(3.0))
    tx_c1 = slide.shapes.add_textbox(Inches(1.0), Inches(3.8), Inches(5.3), Inches(2.6))
    tf_c1 = tx_c1.text_frame
    tf_c1.word_wrap = True
    p = tf_c1.paragraphs[0]
    p.text = "📈 이벤트 유형별 분포 (BarPlot)"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_BLUE_ACCENT
    p_d = tf_c1.add_paragraph()
    p_d.text = "• ZONE_EXIT_DANGER 구역 이탈 건수 추적\n• 데이터프레임 groupby 자동 집계 시각화"
    p_d.font.size = Pt(13)
    p_d.font.color.rgb = COLOR_TEXT_MUTED
    p_d.space_before = Pt(10)

    add_card(slide, Inches(6.8), Inches(3.6), Inches(5.7), Inches(3.0))
    tx_c2 = slide.shapes.add_textbox(Inches(7.0), Inches(3.8), Inches(5.3), Inches(2.6))
    tf_c2 = tx_c2.text_frame
    tf_c2.word_wrap = True
    p = tf_c2.paragraphs[0]
    p.text = "🚨 심각도(Severity)별 분석 (BarPlot)"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_RED
    p_d2 = tf_c2.add_paragraph()
    p_d2.text = "• HIGH vs NORMAL 위험 등급 분류\n• 시각별 안전 위협 빈도 패턴 분석 지원"
    p_d2.font.size = Pt(13)
    p_d2.font.color.rgb = COLOR_TEXT_MUTED
    p_d2.space_before = Pt(10)


def build_slide_8(prs):
    """Slide 08: 이벤트 로그 검색 및 스마트 관제 알림함"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_header(slide, "이벤트 히스토리 검색 & 데이터 검색 기능 (Gradio Dataframe)")

    # 검색 및 데이터 테이블 기능 소개
    add_card(slide, Inches(0.8), Inches(1.8), Inches(11.7), Inches(4.8))
    tx = slide.shapes.add_textbox(Inches(1.1), Inches(2.1), Inches(11.1), Inches(4.2))
    tf = tx.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "🔎 실시간 키워드 필터링 & 데이터 조회 파이프라인"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEXT_MAIN

    bullets = [
        "키워드 대화형 필터: 검색창(kw.change) 입력 시 유형, 카메라, 메시지, 객체명 등 전 컬럼 실시간 라이브 필터링",
        "한글 컬럼 매핑: detected_at(시각), camera_id(카메라), track_id(Track), object(객체), zone(구역), severity(심각도), confidence(신뢰도)",
        "최근 50개 이벤트 역순 정렬(Table Limit 50)을 통해 최신 사고 기록 상단 노출",
        "관제센터 알림함 전문(inbox_full) 동시 조회 지원으로 사고 경위 복원 용이"
    ]
    for b in bullets:
        p_b = tf.add_paragraph()
        p_b.text = "• " + b
        p_b.font.size = Pt(14)
        p_b.font.color.rgb = COLOR_TEXT_MUTED
        p_b.space_before = Pt(14)


def build_slide_9(prs):
    """Slide 09: 핵심 기술 스택 & 다중 카메라 확장성"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_header(slide, "SafeZoo 기술 스택 & 다중 카메라 확장 로드맵")

    # 좌측: 기술 스택
    add_card(slide, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
    tx = slide.shapes.add_textbox(Inches(1.1), Inches(2.1), Inches(5.0), Inches(4.2))
    tf = tx.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "🛠️ Core Tech Stack"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_BLUE_ACCENT

    techs = [
        "Python 3.10+ / OpenCV (비전 처리)",
        "Ultralytics YOLOv8 / YOLOv11 (객체 탐지/추적)",
        "Gradio 5.x (웹 관제 대시보드 UI)",
        "Pandas & NumPy (이벤트 통계 분석)",
        "Custom CSS (다크 & 레드 관제 UI 스타일링)"
    ]
    for t in techs:
        p_t = tf.add_paragraph()
        p_t.text = "• " + t
        p_t.font.size = Pt(13)
        p_t.font.color.rgb = COLOR_TEXT_MAIN
        p_t.space_before = Pt(12)

    # 우측: 멀티 카메라 확장성
    add_card(slide, Inches(6.9), Inches(1.8), Inches(5.6), Inches(4.8))
    tx2 = slide.shapes.add_textbox(Inches(7.2), Inches(2.1), Inches(5.0), Inches(4.2))
    tf2 = tx2.text_frame
    tf2.word_wrap = True
    p2 = tf2.paragraphs[0]
    p2.text = "🎥 Multi-Camera Scalability"
    p2.font.size = Pt(18)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_GREEN

    cams = [
        "CAMERA 01: 사자 방사장 (동물 경계 이탈 감지)",
        "CAMERA 02: 기린 관람 구역 (SAFE/DANGER ROI)",
        "CAMERA 03: 펭귄 관람 구역 (관람객 위험행동 감지)",
        "CAMERA 04: 맹수사 출입 통로 (사람·동물 접근 확인)"
    ]
    for c in cams:
        p_c = tf2.add_paragraph()
        p_c.text = "• " + c
        p_c.font.size = Pt(13)
        p_c.font.color.rgb = COLOR_TEXT_MUTED
        p_c.space_before = Pt(12)


def build_slide_10(prs):
    """Slide 10: 결론 및 Q&A"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)

    # 메인 마무리 카드
    add_card(slide, Inches(1.5), Inches(1.5), Inches(10.3), Inches(4.5))
    tx = slide.shapes.add_textbox(Inches(1.8), Inches(2.0), Inches(9.7), Inches(3.5))
    tf = tx.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "SafeZoo: 모든 생명을 위한 안전한 시선"
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEXT_MAIN
    p.alignment = PP_ALIGN.CENTER

    p_sub = tf.add_paragraph()
    p_sub.text = "24/7 AI 지능형 모니터링으로 안전한 동물원 관제 환경을 선도합니다."
    p_sub.font.size = Pt(18)
    p_sub.font.color.rgb = COLOR_RED
    p_sub.alignment = PP_ALIGN.CENTER
    p_sub.space_before = Pt(16)

    p_q = tf.add_paragraph()
    p_q.text = "\nThank You!  |  Q & A"
    p_q.font.size = Pt(24)
    p_q.font.bold = True
    p_q.font.color.rgb = COLOR_TEXT_MUTED
    p_q.alignment = PP_ALIGN.CENTER
    p_q.space_before = Pt(24)


# ---------------------------------------------------------------------------
# 메인 실행 함수
# ---------------------------------------------------------------------------
def main():
    print("🚀 SafeZoo PPT 자동 생성 에이전트 시작...")
    prs = Presentation()
    prs.slide_width = Inches(13.333) # 16:9 widescreen
    prs.slide_height = Inches(7.5)

    build_slide_1(prs)
    build_slide_2(prs)
    build_slide_3(prs)
    build_slide_4(prs)
    build_slide_5(prs)
    build_slide_6(prs)
    build_slide_7(prs)
    build_slide_8(prs)
    build_slide_9(prs)
    build_slide_10(prs)

    prs.save(OUTPUT_PPTX)
    print(f"✅ SafeZoo PPT 생성 완료: {OUTPUT_PPTX}")


if __name__ == "__main__":
    main()
